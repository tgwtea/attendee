import io
from datetime import UTC, date, datetime, timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import update
from test_attendance import add_member, create

from attendee.application.attendance import AttendanceService
from attendee.application.errors import AccessDenied, NotFound
from attendee.application.reports import ReportService
from attendee.application.responses import ResponseService
from attendee.domain.attendance import SessionStatus
from attendee.domain.reports import cell, is_complete, member_name, session_heading
from attendee.domain.responses import ResponseStatus
from attendee.persistence.models import AttendanceSession
from attendee.reporting.workbook import build_workbook, filename

NOW = datetime(2026, 10, 6, 4, tzinfo=UTC)
ADMIN_TG = 1001
SARAH_TG = 2002


async def set_status(factory, session_id, status):
    async with factory.begin() as session:
        await session.execute(
            update(AttendanceSession)
            .where(AttendanceSession.id == session_id)
            .values(status=status.value)
        )


@pytest.fixture
async def series(attendance_club, session_factory):
    """Closed 13 Oct, Open 20 Oct (Zed joined after 13 Oct), and a Draft 27 Oct."""
    org, admin, _, _ = attendance_club
    await add_member(session_factory, org.id, "Sarah", telegram_id=SARAH_TG)
    first = await create(attendance_club)
    await add_member(session_factory, org.id, "zed")
    second = await create(
        attendance_club,
        series_id=first.series_id,
        new_series_name=None,
        session_date=date(2026, 10, 20),
        deadline=datetime(2026, 10, 19, 12, tzinfo=UTC),
    )
    draft = await create(
        attendance_club,
        series_id=first.series_id,
        new_series_name=None,
        session_date=date(2026, 10, 27),
        deadline=datetime(2026, 10, 26, 12, tzinfo=UTC),
        label="Week 3",
    )
    responses = ResponseService(session_factory)
    for saved in (first, second):
        await set_status(session_factory, saved.id, SessionStatus.OPEN)
    await responses.record(org.id, SARAH_TG, first.id, ResponseStatus.COMING, now=NOW)
    await responses.record(org.id, ADMIN_TG, first.id, ResponseStatus.LATE, "Traffic", now=NOW)
    await responses.record(org.id, SARAH_TG, second.id, ResponseStatus.NOT_COMING, "Sick", now=NOW)
    await set_status(session_factory, first.id, SessionStatus.CLOSED)
    return org, admin, first, second, draft


@pytest.fixture
def service(session_factory):
    return ReportService(session_factory)


def test_cell_values():
    assert cell(True, ResponseStatus.COMING) == 1
    assert cell(True, ResponseStatus.LATE) == 0
    assert cell(True, None) == "NR"
    assert cell(False, None) == "NA"


def test_complete_sessions(archive_after):
    deadline = NOW
    archived = deadline + archive_after + timedelta(seconds=1)
    assert is_complete(SessionStatus.CLOSED, deadline, NOW)
    assert not is_complete(SessionStatus.OPEN, deadline, NOW)
    assert is_complete(SessionStatus.OPEN, deadline, archived)
    assert not is_complete(SessionStatus.DRAFT, deadline, archived)


def test_headings_and_names():
    assert session_heading(date(2026, 10, 7), None) == "7 Oct"
    assert session_heading(date(2026, 10, 7), "Tech Check") == "7 Oct (Tech Check)"
    assert member_name(5, "Ann", "ann") == "Ann"
    assert member_name(5, None, "ann") == "@ann"
    assert member_name(5, None, None) == "Member 5"
    assert filename('A/B: "C"') == "A-B- -C-.xlsx"


async def test_series_report_cells(series, service):
    org, admin, *_ = series
    report = await service.series_report(org.id, admin.id, series[2].series_id, NOW)
    assert report.series_name == "Patrons Day"
    # Draft is left out. Closed comes first and counts; Open follows.
    assert [(column.session_date.day, column.complete) for column in report.columns] == [
        (13, True),
        (20, False),
    ]
    assert [(row.name, row.cells) for row in report.rows] == [
        ("Admin", (0, "NR")),
        ("Member", ("NR", "NR")),
        ("Sarah", (1, 0)),
        ("zed", ("NA", "NR")),
    ]


async def test_archived_open_session_counts_and_leaves_the_list(series, service, archive_after):
    org, admin, first, second, _ = series
    later = second.deadline + archive_after + timedelta(seconds=1)
    report = await service.series_report(org.id, admin.id, first.series_id, later)
    assert [column.complete for column in report.columns] == [True, True]
    _, sessions = await service.list_sessions(org.id, admin.id, first.series_id, later)
    assert sessions == []


async def test_session_report(series, service):
    org, admin, first, *_ = series
    name, sessions = await service.list_sessions(org.id, admin.id, first.series_id, NOW)
    assert name == "Patrons Day"
    assert [session.session_date.day for session in sessions] == [20, 13]
    assert [session.status for session in sessions] == ["Open", "Closed"]
    report = await service.session_report(org.id, admin.id, first.id, NOW)
    assert report.responded == 2 and len(report.members) == 3
    assert report.count(ResponseStatus.LATE) == 1
    assert [(m.name, m.status, m.reason) for m in report.members] == [
        ("Admin", ResponseStatus.LATE, "Traffic"),
        ("Member", None, None),
        ("Sarah", ResponseStatus.COMING, None),
    ]


async def test_reports_need_admin_and_skip_drafts(series, service, attendance_club):
    org, admin, first, _, draft = series
    member = attendance_club[2]
    with pytest.raises(AccessDenied):
        await service.series_report(org.id, member.id, first.series_id, NOW)
    with pytest.raises(AccessDenied):
        await service.session_report(org.id, member.id, first.id, NOW)
    with pytest.raises(NotFound):
        await service.session_report(org.id, admin.id, draft.id, NOW)
    with pytest.raises(NotFound):
        await service.series_report(org.id, admin.id, 999, NOW)
    assert [choice.name for choice in await service.list_series(org.id, admin.id)] == [
        "Patrons Day"
    ]


async def test_draft_only_series_is_not_listed(attendance_club, service):
    org, admin, _, _ = attendance_club
    await create(attendance_club)
    assert await service.list_series(org.id, admin.id) == []


def values(sheet):
    return [[cell.value for cell in row] for row in sheet.iter_rows()]


async def test_workbook_has_live_formulas_over_complete_columns(series, service):
    org, admin, first, *_ = series
    report = await service.series_report(org.id, admin.id, first.series_id, NOW)
    sheet = load_workbook(io.BytesIO(build_workbook(report))).active
    assert sheet is not None
    assert sheet.title == "Patrons Day"
    total = 'C2+D2+COUNTIF(F2:F2,"NR")'
    assert values(sheet)[:2] == [
        ["SN", "Name", "Present", "Absent", "Percentage", "13 Oct", "20 Oct (Open)"],
        [
            1,
            "Admin",
            "=COUNTIF(F2:F2,1)",
            "=COUNTIF(F2:F2,0)",
            f'=IF({total}=0,"",C2/({total}))',
            0,
            "NR",
        ],
    ]
    assert values(sheet)[4][1:2] + values(sheet)[4][5:] == ["zed", "NA", "NR"]
    assert sheet["E2"].number_format == "0.0%"


async def test_workbook_without_complete_sessions_has_zero_totals(series, service):
    org, admin, first, second, _ = series
    await set_status(service.session_factory, first.id, SessionStatus.OPEN)
    report = await service.series_report(org.id, admin.id, first.series_id, NOW)
    sheet = load_workbook(io.BytesIO(build_workbook(report))).active
    assert sheet is not None
    assert values(sheet)[1][2:5] == [0, 0, None]
    assert values(sheet)[0][5:] == ["13 Oct (Open)", "20 Oct (Open)"]


async def test_workbook_writes_formula_like_names_as_text(series, service, session_factory):
    org, admin, first, *_ = series
    await add_member(session_factory, org.id, '=HYPERLINK("http://x","y")')
    later = await create(
        (org, admin, None, AttendanceService(session_factory)),
        series_id=first.series_id,
        new_series_name=None,
        session_date=date(2026, 11, 3),
        deadline=datetime(2026, 11, 2, 12, tzinfo=UTC),
    )
    await set_status(session_factory, later.id, SessionStatus.OPEN)
    report = await service.series_report(org.id, admin.id, first.series_id, NOW)
    assert report.rows[0].name.startswith("=")
    sheet = load_workbook(io.BytesIO(build_workbook(report))).active
    assert sheet is not None
    assert sheet["B2"].data_type == "s"
    assert sheet["B2"].value == '=HYPERLINK("http://x","y")'
