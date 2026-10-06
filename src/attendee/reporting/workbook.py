"""XLSX export of one attendance series (PRD §9, §28). The file exists in memory only (T16).

Present, Absent, and Percentage are live formulas over the complete session columns, so an
edit in the downloaded file updates them (decision T72). Open session columns come after
that range and do not count (decision T73).
"""

import io
import re

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from attendee import copy
from attendee.application.reports import SeriesReport
from attendee.domain.reports import NO_RESPONSE, session_heading

FIXED_HEADINGS = ("SN", "Name", "Present", "Absent", "Percentage")
FIRST_SESSION_COLUMN = len(FIXED_HEADINGS) + 1
_SHEET_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")
_FILE_FORBIDDEN = re.compile(r'[\\/:*?"<>|\x00-\x1f]')


def filename(series_name: str) -> str:
    """'<series>.xlsx' (PRD §28), with characters that file systems reject replaced."""
    return f"{_FILE_FORBIDDEN.sub('-', series_name).strip(' .') or 'attendance'}.xlsx"


def _text(sheet: Worksheet, row: int, column: int, value: str) -> None:
    """Write a string as text. openpyxl stores a string that starts with '=' as a formula."""
    target = sheet.cell(row=row, column=column, value=ILLEGAL_CHARACTERS_RE.sub("", value))
    target.data_type = "s"


def _formulas(sheet: Worksheet, row: int, complete: int) -> None:
    values: tuple[int | str, int | str, str] = (0, 0, "")
    if complete:
        first = get_column_letter(FIRST_SESSION_COLUMN)
        last = get_column_letter(FIRST_SESSION_COLUMN + complete - 1)
        cells = f"{first}{row}:{last}{row}"
        # NA (not on the roster) is outside the denominator; NR (no response) is in it (PRD §9).
        total = f'C{row}+D{row}+COUNTIF({cells},"{NO_RESPONSE}")'
        values = (
            f"=COUNTIF({cells},1)",
            f"=COUNTIF({cells},0)",
            f'=IF({total}=0,"",C{row}/({total}))',
        )
    for column, value in enumerate(values, start=3):
        sheet.cell(row=row, column=column, value=value)
    sheet.cell(row=row, column=5).number_format = "0.0%"


def build_workbook(report: SeriesReport) -> bytes:
    """Return the .xlsx bytes. CPU work: run it in a worker thread."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = (
        _SHEET_FORBIDDEN.sub("-", ILLEGAL_CHARACTERS_RE.sub("", report.series_name))[:31]
        or "Attendance"
    )
    for column, heading in enumerate(FIXED_HEADINGS, start=1):
        sheet.cell(row=1, column=column, value=heading)
    for offset, column in enumerate(report.columns):
        heading = session_heading(column.session_date, column.label)
        if not column.complete:
            heading = copy.STATS_OPEN_COLUMN.format(session=heading)
        _text(sheet, 1, FIRST_SESSION_COLUMN + offset, heading)
    complete = sum(column.complete for column in report.columns)
    for index, member in enumerate(report.rows, start=1):
        row = index + 1
        sheet.cell(row=row, column=1, value=index)
        _text(sheet, row, 2, member.name)
        _formulas(sheet, row, complete)
        for offset, value in enumerate(member.cells):
            sheet.cell(row=row, column=FIRST_SESSION_COLUMN + offset, value=value)
    sheet.freeze_panes = "C2"
    sheet.column_dimensions["B"].width = 28
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
