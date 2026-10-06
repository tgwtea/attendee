"""Read an uploaded namelist from memory. The caller never stores the file."""

import csv
import io
import zipfile
from collections.abc import Iterable
from datetime import date, datetime, time
from pathlib import PurePath

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from attendee import copy
from attendee.domain.imports import ImportFileError, ParsedFile, RawTable, read_table

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 10 * 1024 * 1024
MAX_MEMBER_ROWS = 1000
MAX_COLUMNS = 20
SUFFIXES = (".csv", ".xlsx")


def parse_file(filename: str, data: bytes) -> ParsedFile:
    """Choose the parser from the file name suffix."""
    suffix = PurePath(filename).suffix.lower()
    if suffix == ".csv":
        return parse_csv(data)
    if suffix == ".xlsx":
        return parse_xlsx(data)
    raise ImportFileError(copy.FILE_WRONG_TYPE.format(suffixes=" or ".join(SUFFIXES)))


def parse_csv(data: bytes) -> ParsedFile:
    _check_size(data)
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ImportFileError(copy.FILE_NOT_UTF8) from exc
    try:
        lines = _limited_rows(csv.reader(io.StringIO(text, newline="")))
    except csv.Error as exc:
        raise ImportFileError(copy.FILE_BAD_CSV.format(detail=exc)) from exc
    return read_table(_table(lines))


def parse_xlsx(data: bytes) -> ParsedFile:
    """Read the first worksheet. Formulas give their last saved value."""
    _check_size(data)
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise ImportFileError(copy.FILE_EXPANDED_TOO_LARGE)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ImportFileError(copy.FILE_BAD_XLSX) from exc
    try:
        sheet = workbook.worksheets[0]
        if (sheet.max_row or 0) > MAX_MEMBER_ROWS + 1 or (sheet.max_column or 0) > MAX_COLUMNS:
            raise ImportFileError(copy.FILE_TABLE_TOO_LARGE)
        # Do not trust dimensions that hide cells outside the declared range.
        sheet.reset_dimensions()
        lines = _limited_rows(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    return read_table(_table(lines))


def _limited_rows(rows: Iterable[Iterable[object]]) -> list[list[str]]:
    lines: list[list[str]] = []
    for index, row in enumerate(rows):
        if index > MAX_MEMBER_ROWS:
            raise ImportFileError(copy.FILE_TABLE_TOO_LARGE)
        cells: list[str] = []
        for column, value in enumerate(row):
            if column >= MAX_COLUMNS:
                raise ImportFileError(copy.FILE_TABLE_TOO_LARGE)
            cells.append(_cell_text(value))
        lines.append(cells)
    return lines


def _check_size(data: bytes) -> None:
    if len(data) > MAX_FILE_BYTES:
        raise ImportFileError(copy.FILE_TOO_LARGE.format(megabytes=MAX_FILE_BYTES // (1024 * 1024)))


def _table(lines: list[list[str]]) -> RawTable:
    if not lines:
        raise ImportFileError(copy.FILE_EMPTY)
    return RawTable(tuple(lines[0]), tuple(tuple(line) for line in lines[1:]))


def _cell_text(value: object) -> str:
    """Give XLSX cells the text a CSV export would hold. A whole float loses its ".0"."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).upper()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    return str(value)
